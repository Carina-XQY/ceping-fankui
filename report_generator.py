#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
个性化测评反馈报告生成器
=========================
使用方法：
  1. 准备 Excel 数据文件（表头格式：姓名, 板块名（满分）, ... ）
  2. 编辑 config.json 填写班级、测试名称、教师等信息
  3. 运行: python3 report_generator.py

Excel 表头示例：
  姓名 | 单词（30） | 短语（40） | 单选（15） | 填空（15） | 翻译（20） | 总分
  （括号内为满分，脚本会自动解析）
"""

import base64
import io
import json
import os
import re
import sys
from datetime import datetime

import openpyxl
from PIL import Image
from playwright.sync_api import sync_playwright

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# 板块未在 config.json 中配置时的默认建议
DEFAULT_ADVICE = {
    "topic": "该板块知识",
    "advice": "针对薄弱知识点做专项练习，每天坚持练几道题找找感觉。",
}


# ============================================================
#  图片处理
# ============================================================
def remove_white_bg(path):
    """去除图片白色背景，返回透明 PNG 的 base64"""
    img = Image.open(path).convert("RGBA")
    data = img.getdata()
    new_data = []
    for r, g, b, a in data:
        if r > 235 and g > 235 and b > 235:
            new_data.append((r, g, b, 0))
        else:
            new_data.append((r, g, b, a))
    img.putdata(new_data)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()


def img_to_b64(path):
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode()


# ============================================================
#  Excel 读取（自动解析表头中的满分）
# ============================================================
def load_excel(path):
    """
    返回: (students, sections)
      students = [ (姓名, {板块: 得分, ...}, 总分), ... ]
      sections  = [ (板块名, 满分), ... ]
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) if h else "" for h in rows[0]]

    # 解析表头：提取板块名和满分
    sections = []
    score_col_indices = []
    total_col_idx = None
    name_col_idx = None

    for i, h in enumerate(headers):
        if h == "姓名":
            name_col_idx = i
            continue
        if h in ("总分", "合计", "Total"):
            total_col_idx = i
            continue
        # 尝试从 "单词（30）" 格式中提取满分
        m = re.search(r"[（(](\d+(?:\.\d+)?)[）)]", h)
        if m:
            name = re.sub(r"[（(].*?[）)]", "", h).strip()
            max_score = float(m.group(1))
            sections.append((name, max_score))
            score_col_indices.append(i)

    students = []
    for row in rows[1:]:
        name = row[name_col_idx]
        if not name:
            continue
        scores = {}
        for sec_idx, col_idx in enumerate(score_col_indices):
            sec_name, sec_max = sections[sec_idx]
            val = row[col_idx]
            scores[sec_name] = float(val) if val is not None else 0

        if total_col_idx is not None and row[total_col_idx] is not None:
            total = float(row[total_col_idx])
        else:
            total = sum(scores.values())

        students.append((str(name), scores, total))

    return students, sections


# ============================================================
#  个性化内容生成
# ============================================================
def gen_overall_eval(name, total, total_max, scores, sections):
    """综合评价：按总分率分档"""
    rate = total / total_max
    # 找最强和最弱板块
    rates = {s: scores[s] / m for s, m in sections if s in scores}
    strongest = max(rates, key=rates.get) if rates else ""
    weakest = min(rates, key=rates.get) if rates else ""
    word_rate = rates.get("单词", 0)

    if rate >= 0.75:
        text = (f"本次{config['test_name']}取得 <b>{total}</b> 分，整体表现优异。"
                f"{strongest}板块掌握扎实，说明平时有认真记背，值得表扬。"
                f"{'翻译板块表现突出，句子组织能力较强。' if rates.get('翻译', 0) >= 0.7 else ''}"
                f"语法运用板块仍有提升空间，补上来分数还能再上一个台阶。")
    elif rate >= 0.65:
        text = (f"本次{config['test_name']}取得 <b>{total}</b> 分，整体表现良好。"
                f"{'词汇基础扎实，短语搭配记得牢。' if word_rate >= 0.85 else '词汇短语基础尚可，个别不太熟的需要巩固。'}"
                f"{weakest}板块失分较多，这是需要重点突破的方向。"
                f"孩子底子是有的，薄弱板块补上来分数还能再上一个台阶。")
    elif rate >= 0.55:
        text = (f"本次{config['test_name']}取得 <b>{total}</b> 分，整体表现中等。"
                f"{'词汇基础尚可，但短语搭配还需加强。' if word_rate >= 0.7 else '词汇基础还需巩固，部分知识点记得不够牢。'}"
                f"{weakest}板块是主要失分点，需要多做练习找找感觉。"
                f"孩子其实不笨，就是有些基础没跟上，咱们一步步来，肯定能提上来的。")
    else:
        text = (f"本次{config['test_name']}取得 <b>{total}</b> 分，成绩不太理想，"
                f"主要原因是{'词汇基础薄弱' if word_rate < 0.7 else '基础不够扎实'}影响了整体发挥。"
                f"{weakest}板块也因基础受限而失分较多。"
                f"孩子现在最需要的是先把基础打牢，基础上来了其他板块也会跟着好起来。"
                f"别着急，咱们一起帮孩子把基础补上。")
    return text


def gen_badges(scores, sections):
    """学习亮点：按各板块得分率颁发徽章"""
    badges = []
    for sec_name, max_score in sections:
        if sec_name not in scores:
            continue
        rate = scores[sec_name] / max_score
        if rate >= 1.0:
            badges.append(("🌟", f"{sec_name}满分"))
        elif rate >= 0.9:
            badges.append(("📚", f"{sec_name}达人"))
    if not badges:
        badges.append(("💪", "坚持之星"))
    return badges


def gen_suggestions(scores, sections, section_advice):
    """成长建议：取得分率最低的3个板块"""
    ranked = sorted(sections, key=lambda x: scores.get(x[0], 0) / x[1])
    suggestions = []
    for sec_name, _ in ranked[:3]:
        advice_info = section_advice.get(sec_name, DEFAULT_ADVICE)
        suggestions.append((advice_info["topic"], advice_info["advice"]))
    return suggestions


def gen_teacher_msg(name, total, total_max):
    """教师寄语：按分数段配英文格言"""
    rate = total / total_max
    if rate >= 0.75:
        msg = f"{name}同学，你的努力大家都看得到，继续保持，薄弱板块再精进一步就更棒了！"
        quote = '"Practice makes perfect!" 熟能生巧，你的坚持终将换来质的飞跃。'
    elif rate >= 0.65:
        msg = f"{name}同学，底子不错，只要薄弱板块补上来，分数还能再上一个台阶，加油！"
        quote = '"Every expert was once a beginner." 每位高手都曾是初学者，坚持你就一定能行。'
    elif rate >= 0.55:
        msg = f"{name}同学，你其实不笨，就是有些基础没跟上，咱们一步步来，肯定能提上来的！"
        quote = '"Slow and steady wins the race." 稳扎稳打，终会成功。'
    elif rate >= 0.45:
        msg = f"{name}同学，基础关先过了，其他板块会跟着好起来。多鼓励自己，咱们一起努力！"
        quote = '"A journey of a thousand miles begins with a single step." 千里之行始于足下。'
    else:
        msg = f"{name}同学，现在最重要的是先把基础打牢，每天坚持一点点，你会看到变化的！"
        quote = '"Fall seven times, stand up eight." 跌倒了就爬起来，坚持就是胜利。'
    return msg, quote


def star_rating(score, max_score):
    """星级评分（5星制）"""
    full = int(score / max_score * 5) if max_score > 0 else 0
    return "★" * full + "☆" * (5 - full)


# ============================================================
#  HTML 模板
# ============================================================
def build_html(name, scores, total, sections, total_max):
    overall = gen_overall_eval(name, total, total_max, scores, sections)
    badges = gen_badges(scores, sections)
    suggestions = gen_suggestions(scores, sections, config["section_advice"])
    teacher_msg, quote = gen_teacher_msg(name, total, total_max)

    # 分项评分行
    score_rows = ""
    for sec_name, max_score in sections:
        s = scores.get(sec_name, 0)
        stars = star_rating(s, max_score)
        score_rows += f"""
        <tr>
            <td class="sec-name">{sec_name}</td>
            <td class="sec-score">{s}<span class="max">/{max_score:g}</span></td>
            <td class="sec-rate">{(s/max_score*100):.0f}%</td>
            <td class="sec-star">{stars}</td>
        </tr>"""

    # 徽章
    badge_html = ""
    for emoji, label in badges:
        badge_html += f'<div class="badge"><span class="badge-emoji">{emoji}</span><span class="badge-text">{label}</span></div>'

    # 建议列表
    sug_html = ""
    for topic, advice in suggestions:
        sug_html += f'<li><span class="sug-topic">{topic}：</span>{advice}</li>'

    date_str = config["date"]
    if date_str == "auto":
        date_str = datetime.now().strftime("%Y年%m月%d日")

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<style>
* {{ margin: 0; padding: 0; box-sizing: border-box; }}
body {{
    font-family: "PingFang SC", "Microsoft YaHei", "Helvetica Neue", sans-serif;
    background: linear-gradient(135deg, #E3F2FD 0%, #E1F5FE 50%, #E8EAF6 100%);
    padding: 30px;
    width: 820px;
}}
.report {{
    background: #FFFFFF;
    border-radius: 24px;
    overflow: hidden;
    box-shadow: 0 8px 32px rgba(33, 150, 243, 0.15);
    position: relative;
}}
.header {{
    background: linear-gradient(135deg, #42A5F5 0%, #1E88E5 50%, #1565C0 100%);
    padding: 28px 36px 24px;
    color: white;
    position: relative;
    overflow: hidden;
}}
.header::before {{ content: "✨"; position: absolute; top: 10px; right: 20px; font-size: 28px; opacity: 0.5; }}
.header::after {{ content: "✨"; position: absolute; bottom: 8px; right: 60px; font-size: 18px; opacity: 0.4; }}
.brand-row {{ display: flex; align-items: center; gap: 10px; margin-bottom: 14px; }}
.brand-logo {{ height: 40px; }}
.brand-name {{ font-size: 14px; font-weight: 600; letter-spacing: 1px; opacity: 0.95; }}
.title-row {{ display: flex; align-items: center; gap: 12px; }}
.title-icon {{ font-size: 32px; }}
.title {{ font-size: 26px; font-weight: 700; }}
.subtitle {{ font-size: 14px; opacity: 0.9; margin-top: 4px; }}
.body {{ padding: 28px 36px 28px; }}
.section {{ margin-bottom: 22px; }}
.sec-title {{ font-size: 17px; font-weight: 700; color: #1565C0; margin-bottom: 12px; display: flex; align-items: center; gap: 8px; }}
.sec-title-icon {{ font-size: 20px; }}
.eval-box {{ background: #E1F5FE; border-left: 4px solid #42A5F5; border-radius: 12px; padding: 16px 20px; font-size: 14px; line-height: 1.8; color: #0D47A1; }}
.score-table {{ width: 100%; border-collapse: collapse; border-radius: 12px; overflow: hidden; box-shadow: 0 0 0 1px #BBDEFB; }}
.score-table th {{ background: #BBDEFB; color: #1565C0; font-size: 13px; padding: 10px 16px; text-align: left; }}
.score-table td {{ padding: 10px 16px; font-size: 15px; border-bottom: 1px solid #E3F2FD; }}
.sec-name {{ color: #1565C0; font-weight: 600; }}
.sec-score {{ font-weight: 700; color: #0D47A1; font-size: 18px; }}
.sec-score .max {{ font-size: 13px; color: #9E9E9E; font-weight: 400; }}
.sec-rate {{ color: #42A5F5; font-size: 14px; }}
.sec-star {{ color: #FFA726; letter-spacing: 2px; }}
.score-table tr:last-child td {{ border-bottom: none; }}
.total-row td {{ background: #E1F5FE; font-weight: 700; border-top: 2px solid #42A5F5; }}
.badges {{ display: flex; flex-wrap: wrap; gap: 12px; }}
.badge {{ background: linear-gradient(135deg, #FFF3E0, #FFE0B2); border: 1.5px solid #FFB74D; border-radius: 20px; padding: 8px 18px; display: flex; align-items: center; gap: 6px; box-shadow: 0 2px 8px rgba(255, 167, 38, 0.2); }}
.badge-emoji {{ font-size: 18px; }}
.badge-text {{ font-size: 14px; font-weight: 600; color: #E65100; }}
.sug-list {{ list-style: none; counter-reset: sug; }}
.sug-list li {{ counter-increment: sug; background: #E1F5FE; border-radius: 10px; padding: 12px 16px 12px 48px; margin-bottom: 10px; font-size: 14px; line-height: 1.7; color: #0D47A1; position: relative; }}
.sug-list li::before {{ content: counter(sug); position: absolute; left: 14px; top: 50%; transform: translateY(-50%); width: 24px; height: 24px; background: #42A5F5; color: white; border-radius: 50%; display: flex; align-items: center; justify-content: center; font-size: 13px; font-weight: 700; }}
.sug-topic {{ color: #1565C0; font-weight: 700; }}
.msg-box {{ background: linear-gradient(135deg, #E3F2FD, #BBDEFB); border-radius: 16px; padding: 20px 24px; }}
.msg-text {{ font-size: 14px; line-height: 1.8; color: #1565C0; margin-bottom: 10px; }}
.msg-quote {{ font-size: 15px; font-style: italic; color: #0D47A1; font-weight: 600; text-align: right; }}
.footer {{ background: #E1F5FE; padding: 16px 36px; display: flex; justify-content: space-between; align-items: center; border-top: 1px solid #BBDEFB; }}
.footer-info {{ font-size: 13px; color: #1565C0; }}
.footer-mascot {{ height: 60px; }}
.heart {{ position: absolute; color: #90CAF9; opacity: 0.3; font-size: 16px; }}
.heart-1 {{ top: 200px; left: 12px; }}
.heart-2 {{ bottom: 100px; right: 12px; font-size: 20px; }}
</style>
</head>
<body>
<div class="report">
    <div class="header">
        <div class="brand-row">
            <img class="brand-logo" src="data:image/png;base64,{LOGO_B64}" alt="logo">
            <span class="brand-name">飞扬精准自学 · AFTER-SCHOOL CARE CENTER</span>
        </div>
        <div class="title-row">
            <span class="title-icon">📖</span>
            <div>
                <div class="title">{name}同学{config['semester']}{config['subject']}学习反馈</div>
                <div class="subtitle">{config['semester']} · {config['test_name']} · 满分{total_max:g}分</div>
            </div>
        </div>
    </div>
    <div class="body">
        <span class="heart heart-1">💙</span>
        <span class="heart heart-2">💙</span>
        <div class="section">
            <div class="sec-title"><span class="sec-title-icon">📋</span>综合评价</div>
            <div class="eval-box">{overall}</div>
        </div>
        <div class="section">
            <div class="sec-title"><span class="sec-title-icon">📊</span>分项评分</div>
            <table class="score-table">
                <tr><th>板块</th><th>得分</th><th>得分率</th><th>星级</th></tr>
                {score_rows}
                <tr class="total-row">
                    <td class="sec-name" style="font-size:16px;">总分</td>
                    <td class="sec-score" style="font-size:20px;">{total:g}<span class="max">/{total_max:g}</span></td>
                    <td class="sec-rate">{(total/total_max*100):.0f}%</td>
                    <td class="sec-star">{star_rating(total, total_max)}</td>
                </tr>
            </table>
        </div>
        <div class="section">
            <div class="sec-title"><span class="sec-title-icon">🏆</span>学习亮点</div>
            <div class="badges">{badge_html}</div>
        </div>
        <div class="section">
            <div class="sec-title"><span class="sec-title-icon">💡</span>成长建议</div>
            <ol class="sug-list">{sug_html}</ol>
        </div>
        <div class="section">
            <div class="sec-title"><span class="sec-title-icon">💌</span>教师寄语</div>
            <div class="msg-box">
                <div class="msg-text">{teacher_msg}</div>
                <div class="msg-quote">{quote}</div>
            </div>
        </div>
    </div>
    <div class="footer">
        <div class="footer-info">教师：{config['teacher']}　|　{date_str}</div>
        <img class="footer-mascot" src="data:image/png;base64,{MASCOT_B64}" alt="mascot">
    </div>
</div>
</body>
</html>"""
    return html


# ============================================================
#  主流程
# ============================================================
def main():
    global config, LOGO_B64, MASCOT_B64

    # 1. 加载配置
    config_path = os.path.join(BASE_DIR, "config.json")
    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    # 2. 加载图片资源
    logo_path = os.path.join(BASE_DIR, config["logo"])
    mascot_path = os.path.join(BASE_DIR, config["mascot"])
    LOGO_B64 = remove_white_bg(logo_path)
    MASCOT_B64 = img_to_b64(mascot_path)

    # 3. 读取 Excel
    data_path = os.path.join(BASE_DIR, config["data_file"])
    students, sections = load_excel(data_path)
    total_max = sum(m for _, m in sections)

    print(f"班级：{config['class_name']}")
    print(f"测试：{config['test_name']}")
    print(f"满分：{total_max:g} 分")
    print(f"板块：{', '.join(f'{n}({m:g})' for n, m in sections)}")
    print(f"人数：{len(students)} 人")
    print("-" * 50)

    # 4. 创建输出目录
    output_dir = os.path.join(BASE_DIR, config.get("output_dir", "反馈报告"))
    os.makedirs(output_dir, exist_ok=True)

    # 5. 逐个生成报告
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 880, "height": 1200})

        for name, scores, total in students:
            html = build_html(name, scores, total, sections, total_max)
            html_path = os.path.join(output_dir, f"{name}_temp.html")
            with open(html_path, "w", encoding="utf-8") as f:
                f.write(html)

            page.goto(f"file://{html_path}")
            page.wait_for_timeout(800)

            report_el = page.query_selector(".report")
            img_path = os.path.join(output_dir, f"{name}_反馈报告.png")
            report_el.screenshot(path=img_path)
            os.remove(html_path)

            rate = total / total_max * 100
            print(f"  ✓ {name}　{total:g}/{total_max:g}（{rate:.0f}%）")

        browser.close()

    # 6. 汇总
    print("-" * 50)
    print(f"完成！{len(students)} 份报告已保存到：{output_dir}")


if __name__ == "__main__":
    main()
